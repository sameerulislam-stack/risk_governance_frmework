from fastapi import FastAPI, Request, UploadFile, File, Form
from fastapi import HTTPException
from fastapi.templating import Jinja2Templates
from fastapi.staticfiles import StaticFiles
from fastapi.responses import RedirectResponse, HTMLResponse, Response
from io import BytesIO
import sqlite3
import os
import sys
import shutil
import math
import re
import tempfile
import zipfile
from datetime import date, datetime
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from etl import run_etl

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
RESOURCE_DIR = getattr(sys, "_MEIPASS", BASE_DIR)
DATA_DIR = os.environ.setdefault(
    "JLI_DATA_DIR",
    os.path.join(BASE_DIR, "data"),
)
os.makedirs(DATA_DIR, exist_ok=True)

app = FastAPI()
templates = Jinja2Templates(directory=os.path.join(RESOURCE_DIR, "templates"))
static_dir = os.path.join(RESOURCE_DIR, "static")
if os.path.isdir(static_dir):
    app.mount("/static", StaticFiles(directory=static_dir), name="static")

def render(request: Request, template_name: str, context: dict = None):
    if context is None: context = {}
    context["request"] = request
    return HTMLResponse(templates.env.get_template(template_name).render(context))

GOVERNANCE_STRUCTURE = {
    "MARKET & LIQUIDITY RISK": [
        "RISK LIMITS (FUND WISE )",
        "TRADE MONITORING",
        "STOP LOSS MONITORING",
        "MONITORING OF PORTFOLIO RISK PARAMETERS (NAV BASED)",
        "MONITORING OF  EQUITY PORTFOLIO RISK PARAMETERS",
        "INVESTMENT INCOME ATTRIBUTION",
        "PORTFOLIO EQUITY ATTRIBUTION"
    ],
    "CREDIT RISK": [
        "DEBT SECURITIES RISK REVIEW",
        "MONITORING OF TFC's PORTFOLIO"
    ],
    "OPERATIONAL RISK": [
        "TRADING ROOM RISK CONTROLS",
        "KEY RISK INDICATORS",
        "RISK REGISTERS"
    ]
}

EQUITY_REPORT_FUNDS = [
    "Consolidated",
    "Jubilee Life Aggressive Fund",
    "Jubilee Life Balanced Fund",
    "Meesaq Balanced Fund",
    "Aggressive Takaful Fund",
    "Balanced Takaful Fund",
    "Family Takaful Balance Fund",
]

def get_db():
    return sqlite3.connect(os.path.join(DATA_DIR, "risk.db"))


TRADE_MONITORING_SHEET = "Trade Risk Analysis"
TRADE_MONITORING_FILENAME = "trade_monitoring.xlsx"
TRADE_MONITORING_SEED_FILENAME = "Trade Monitoring - All - 2026.xlsx"


def _normalize_trade_header(value):
    if value is None or pd.isna(value):
        return ""
    normalized = re.sub(r"\s+", " ", str(value)).strip().casefold()
    return re.sub(r"\s*/\s*", "/", normalized)


def _read_trade_monitoring_workbook(path):
    try:
        workbook = pd.ExcelFile(path)
    except (OSError, ValueError, zipfile.BadZipFile) as error:
        raise ValueError(f"Could not read the Excel workbook: {error}") from error

    try:
        sheet_name = next(
            (
                name
                for name in workbook.sheet_names
                if name.strip().casefold() == TRADE_MONITORING_SHEET.casefold()
            ),
            None,
        )
        if sheet_name is None:
            raise ValueError(
                f"The workbook must contain a sheet titled '{TRADE_MONITORING_SHEET}'."
            )
        raw = pd.read_excel(workbook, sheet_name=sheet_name, header=None)
    finally:
        workbook.close()
    required_headers = {"date", "scrip", "buy/sell"}
    header_row = None
    for row_index in range(min(len(raw), 15)):
        row_headers = {
            _normalize_trade_header(value)
            for value in raw.iloc[row_index].tolist()
        }
        if required_headers.issubset(row_headers):
            header_row = row_index
            break
    if header_row is None:
        raise ValueError(
            f"Could not find the Date, Scrip, and BUY/SELL headers on "
            f"'{TRADE_MONITORING_SHEET}'."
        )

    headers = raw.iloc[header_row].tolist()
    retained_columns = [
        (index, str(value).strip())
        for index, value in enumerate(headers)
        if value is not None and not pd.isna(value) and str(value).strip()
    ]
    columns = [name for _, name in retained_columns]
    normalized_headers = {_normalize_trade_header(name) for name in columns}
    if len(normalized_headers) != len(columns):
        raise ValueError("The Trade Risk Analysis sheet has duplicate column headers.")
    required_report_columns = {
        "date",
        "scrip",
        "buy/sell",
        "amount",
        "amount vested",
        "market participation",
        "stock performance",
        "optimal/sub-optimal",
        "overall score",
    }
    missing_columns = required_report_columns - normalized_headers
    if missing_columns:
        raise ValueError(
            "The Trade Risk Analysis sheet is missing required columns: "
            + ", ".join(sorted(missing_columns))
        )

    data = raw.iloc[header_row + 1 :, [index for index, _ in retained_columns]].copy()
    data.columns = columns
    date_column = next(
        name for name in columns if _normalize_trade_header(name) == "date"
    )
    data = data.loc[data[date_column].notna()].copy()
    data[date_column] = pd.to_datetime(data[date_column], errors="coerce")
    if data[date_column].isna().any():
        raise ValueError("The Trade Risk Analysis sheet contains invalid trade dates.")
    data = data.reset_index(drop=True)
    return data


def _trade_column_map(data):
    return {
        _normalize_trade_header(name): name
        for name in data.columns
    }


def _trade_numeric_column(data, column_map, name):
    column = column_map.get(_normalize_trade_header(name))
    if column is None:
        return pd.Series(0.0, index=data.index)
    return pd.to_numeric(data[column], errors="coerce").fillna(0.0)


def _trade_summary(data):
    columns = _trade_column_map(data)
    count = len(data)
    amount_column = columns.get(_normalize_trade_header("Amount Vested"))
    if amount_column is None:
        amount_column = columns.get(_normalize_trade_header("Amount"))
    values = (
        pd.to_numeric(data[amount_column], errors="coerce").fillna(0.0).abs()
        if amount_column is not None
        else pd.Series(0.0, index=data.index)
    )

    side_column = columns.get(_normalize_trade_header("BUY/SELL"))
    sides = (
        data[side_column].astype(str).str.strip().str.upper()
        if side_column is not None
        else pd.Series("", index=data.index)
    )
    buy_mask = sides.str.startswith("BUY")
    sell_mask = sides.str.startswith("SELL")
    total_value = float(values.sum())

    def grouped_rows(labels, masks):
        output = []
        for label, mask in zip(labels, masks):
            row_count = int(mask.sum())
            row_value = float(values.loc[mask].sum())
            output.append(
                {
                    "label": label,
                    "value": row_value,
                    "count": row_count,
                    "value_percent": row_value / total_value if total_value else None,
                    "count_percent": row_count / count if count else None,
                }
            )
        return output

    optimal_column = columns.get(_normalize_trade_header("Optimal/Sub-optimal"))
    if optimal_column is not None:
        optimal_labels = data[optimal_column].fillna("Unclassified").astype(str).str.strip()
        optimal_normalized = optimal_labels.str.casefold()
    else:
        optimal_labels = pd.Series("Unclassified", index=data.index)
        optimal_normalized = optimal_labels.str.casefold()
    optimal_mask = optimal_normalized.str.startswith("optimal") & ~optimal_normalized.str.startswith("sub")
    suboptimal_mask = optimal_normalized.str.startswith("sub")
    other_optimal_mask = ~(optimal_mask | suboptimal_mask)

    performance_column = columns.get(_normalize_trade_header("Stock Performance"))
    performance_labels = (
        data[performance_column].fillna("Unclassified").astype(str).str.strip()
        if performance_column is not None
        else pd.Series("Unclassified", index=data.index)
    )
    actual_performance_labels = list(dict.fromkeys(performance_labels.tolist()))
    preferred_performance_labels = ["Highly Liquid", "Liquid", "Illiquid"]
    ordered_performance_labels = [
        next(
            label
            for label in actual_performance_labels
            if label.casefold() == preferred.casefold()
        )
        for preferred in preferred_performance_labels
        if any(
            label.casefold() == preferred.casefold()
            for label in actual_performance_labels
        )
    ]
    ordered_performance_labels.extend(
        label
        for label in actual_performance_labels
        if label.casefold()
        not in {preferred.casefold() for preferred in preferred_performance_labels}
    )
    performance_groups = [
        (label, performance_labels.str.casefold() == label.casefold())
        for label in ordered_performance_labels
    ]

    participation = _trade_numeric_column(data, columns, "Market Participation")
    if participation.abs().max() > 1:
        participation = participation / 100.0
    participation_groups = [
        ("30%+", participation >= 0.30),
        ("15%-30%", (participation >= 0.15) & (participation < 0.30)),
        ("Less than 15%", participation < 0.15),
    ]

    score_column = columns.get(_normalize_trade_header("Overall Score"))
    if score_column is not None:
        scores = pd.to_numeric(data[score_column], errors="coerce")
        available_scores = sorted(
            set(range(2, 9)).union(scores.dropna().astype(int).tolist())
        )
        score_distribution = [
            {"score": str(score), "count": int((scores == score).sum())}
            for score in available_scores
        ]
    else:
        score_distribution = []

    return {
        "total_executions": count,
        "total_value_vested": total_value,
        "buy_executions": int(buy_mask.sum()),
        "buy_value": float(values.loc[buy_mask].sum()),
        "sell_executions": int(sell_mask.sum()),
        "sell_value": float(values.loc[sell_mask].sum()),
        "optimal_rate": float(values.loc[optimal_mask].sum()) / total_value if total_value else None,
        "suboptimal_rate": float(values.loc[suboptimal_mask].sum()) / total_value if total_value else None,
        "unclassified_optimal_value": float(values.loc[other_optimal_mask].sum()),
        "optimal_summary": grouped_rows(
            ["Optimal Trade", "Sub-optimal Trade", "Unclassified"],
            [optimal_mask, suboptimal_mask, other_optimal_mask],
        ),
        "stock_performance": grouped_rows(
            [label for label, _ in performance_groups],
            [mask for _, mask in performance_groups],
        ),
        "market_participation": grouped_rows(
            [label for label, _ in participation_groups],
            [mask for _, mask in participation_groups],
        ),
        "score_distribution": score_distribution,
    }


def _trade_json_value(value):
    if value is None or pd.isna(value):
        return None
    if isinstance(value, (pd.Timestamp, datetime, date)):
        return value.isoformat()[:10]
    if hasattr(value, "item"):
        value = value.item()
    if isinstance(value, float) and not math.isfinite(value):
        return None
    return value


def _trade_workbook_path():
    uploaded_path = os.path.join(DATA_DIR, TRADE_MONITORING_FILENAME)
    if os.path.isfile(uploaded_path):
        return uploaded_path
    seed_path = os.path.join(RESOURCE_DIR, "images", TRADE_MONITORING_SEED_FILENAME)
    if os.path.isfile(seed_path):
        return seed_path
    raise HTTPException(
        status_code=503,
        detail="Trade Monitoring data is unavailable. Upload a workbook containing the Trade Risk Analysis sheet.",
    )


def get_equity_funds():
    conn = get_db()
    try:
        funds = pd.read_sql_query(
            "SELECT DISTINCT fund_name FROM fact_portfolio_history "
            "ORDER BY fund_name",
            conn,
        )["fund_name"].dropna().tolist()
    finally:
        conn.close()
    return sorted(funds, key=lambda fund: (fund != "Consolidated", fund.casefold()))

@app.get("/")
async def root():
    return RedirectResponse(url="/dashboard.html")

@app.get("/dashboard.html")
async def get_dashboard(request: Request):
    return render(request, "dashboard.html", {"structure": GOVERNANCE_STRUCTURE})

@app.get("/upload")
async def get_upload_page(request: Request):
    return render(request, "upload.html")


@app.post("/api/upload/trade-monitoring")
async def handle_trade_monitoring_upload(trade_workbook: UploadFile = File(...)):
    os.makedirs(DATA_DIR, exist_ok=True)
    final_path = os.path.join(DATA_DIR, TRADE_MONITORING_FILENAME)
    temporary_path = None

    try:
        with tempfile.NamedTemporaryFile(
            mode="wb", suffix=".xlsx", dir=DATA_DIR, delete=False
        ) as temporary_file:
            temporary_path = temporary_file.name
            shutil.copyfileobj(trade_workbook.file, temporary_file)
        _read_trade_monitoring_workbook(temporary_path)
        os.replace(temporary_path, final_path)
        temporary_path = None
    except (OSError, ValueError, zipfile.BadZipFile) as error:
        raise HTTPException(
            status_code=422,
            detail=f"Trade Monitoring workbook upload failed: {error}",
        ) from error
    finally:
        if temporary_path and os.path.exists(temporary_path):
            os.remove(temporary_path)

    return RedirectResponse(url="/report/trade-monitoring", status_code=303)


@app.post("/api/upload")
async def handle_upload(
    w1: UploadFile = File(...), w2: UploadFile = File(...), w3: UploadFile = File(...),
    data_date: str = Form(...)
):
    os.makedirs(DATA_DIR, exist_ok=True)
    paths = {}
    for key, upload in (("w1", w1), ("w2", w2), ("w3", w3)):
        original_name = (upload.filename or f"{key}.xlsx").replace("\\", "/")
        safe_name = os.path.basename(original_name) or f"{key}.xlsx"
        paths[key] = os.path.join(DATA_DIR, f"{key}_{safe_name}")
        with open(paths[key], "wb") as workbook_file:
            shutil.copyfileobj(upload.file, workbook_file)
    run_etl(paths["w1"], paths["w2"], paths["w3"])
    return RedirectResponse(url="/dashboard.html", status_code=303)

# DYNAMIC REPORT ROUTER
@app.get("/report/{category}/{sub_category}")
async def get_dynamic_report(request: Request, category: str, sub_category: str):
    # Map URL slugs to specific templates
    slug = sub_category.lower()

    if "trade-monitoring" in slug:
        return render(request, "Report_trade_monitoring.html")
    elif "equity-risk-parameters" in slug:
        return render(
            request,
            "report_equity.html",
            {"funds": EQUITY_REPORT_FUNDS},
        )
    else:
        return render(
            request,
            "report.html",
            {"category": sub_category.replace("-", " ").upper(), "data": []},
        )


@app.get("/report/trade-monitoring")
async def get_trade_monitoring_report(request: Request):
    return render(request, "Report_trade_monitoring.html")


@app.get("/report/{category}")
async def get_placeholder_report(request: Request, category: str):
    return render(
        request,
        "report.html",
        {"category": category.replace("-", " ").upper(), "data": []},
    )


@app.get("/api/trade-monitoring")
async def get_trade_monitoring_data(
    mode: str = "imc",
    from_date: str | None = None,
    to_date: str | None = None,
):
    if mode not in {"imc", "cy-2026"}:
        raise HTTPException(status_code=422, detail="Mode must be 'imc' or 'cy-2026'.")

    try:
        data = _read_trade_monitoring_workbook(_trade_workbook_path())
    except ValueError as error:
        raise HTTPException(
            status_code=503,
            detail=f"Trade Monitoring data could not be loaded: {error}",
        ) from error

    columns = _trade_column_map(data)
    date_column = columns["date"]
    if mode == "cy-2026":
        selected = data.loc[
            (data[date_column] >= pd.Timestamp("2026-01-01"))
            & (data[date_column] < pd.Timestamp("2027-01-01"))
        ].copy()
        period_start = "2026-01-01"
        period_end = "2026-12-31"
    else:
        try:
            start = date.fromisoformat(from_date) if from_date else data[date_column].min().date()
            end = date.fromisoformat(to_date) if to_date else data[date_column].max().date()
        except ValueError as error:
            raise HTTPException(
                status_code=422,
                detail="From and To dates must use YYYY-MM-DD format.",
            ) from error
        if start > end:
            raise HTTPException(status_code=422, detail="From date must not be after To date.")
        selected = data.loc[
            data[date_column].between(pd.Timestamp(start), pd.Timestamp(end))
        ].copy()
        period_start = start.isoformat()
        period_end = end.isoformat()

    rows = [
        {
            str(header): _trade_json_value(value)
            for header, value in row.items()
        }
        for row in selected.to_dict(orient="records")
    ]
    return {
        "mode": mode,
        "from_date": period_start,
        "to_date": period_end,
        "as_of_date": (
            selected[date_column].max().date().isoformat() if not selected.empty else None
        ),
        "available_from": data[date_column].min().date().isoformat(),
        "available_to": data[date_column].max().date().isoformat(),
        "headers": [str(header) for header in selected.columns],
        "rows": rows,
        "overview": _trade_summary(selected),
    }


@app.get("/api/equity-data/{fund_name}")
async def get_equity_data(fund_name: str):
    conn = get_db()
    try:
        df = pd.read_sql_query(
            "SELECT date, price FROM fact_portfolio_history "
            "WHERE fund_name = ? AND scrip_code = 'PORTFOLIO' ORDER BY date",
            conn,
            params=(fund_name.strip(),),
        )
    finally:
        conn.close()
    return df.to_dict(orient="records")

@app.get("/api/equity-report/{fund_name}")
async def get_equity_report_data(
    fund_name: str,
    period: str = "1d",
    month: str | None = None,
):
    if period not in {"1d", "1w", "1m", "3m"}:
        raise HTTPException(
            status_code=422,
            detail="Period must be one of: 1d, 1w, 1m, 3m.",
        )

    conn = get_db()
    try:
        tables = {
            row[0]
            for row in conn.execute(
                "SELECT name FROM sqlite_master WHERE type = 'table'"
            )
        }
        portfolio = pd.read_sql_query(
            "SELECT date, price FROM fact_portfolio_history "
            "WHERE fund_name = ? AND scrip_code = 'PORTFOLIO' ORDER BY date",
            conn,
            params=(fund_name.strip(),),
            parse_dates=["date"],
        )
        if "fact_scrip_price_history" not in tables:
            raise HTTPException(
                status_code=503,
                detail="W1 price history is unavailable. Upload W1, W2, and W3 first.",
            )
        if "fact_scrip_volume_history" not in tables:
            raise HTTPException(
                status_code=503,
                detail="W1 volume history is unavailable. Reprocess the uploaded W1, W2, and W3 workbooks.",
            )
        market = pd.read_sql_query(
            "SELECT date, price FROM fact_scrip_price_history "
            "WHERE scrip_code = 'KSE30' ORDER BY date",
            conn,
            parse_dates=["date"],
        )
        if "fact_portfolio_weights" not in tables:
            raise HTTPException(
                status_code=503,
                detail="Portfolio holdings are unavailable. Reprocess W1, W2, and W3.",
            )
        weights = pd.read_sql_query(
            "SELECT date, scrip_code, holding, price, market_value, portfolio_weight "
            "FROM fact_portfolio_weights WHERE fund_name = ? ORDER BY scrip_code",
            conn,
            params=(fund_name.strip(),),
        )
        symbols = weights["scrip_code"].dropna().unique().tolist()
        if not symbols:
            raise HTTPException(
                status_code=404,
                detail=f"No holdings are available for '{fund_name}'.",
            )

        placeholders = ",".join("?" for _ in symbols)
        security_prices = pd.read_sql_query(
            "SELECT date, scrip_code, price FROM fact_scrip_price_history "
            f"WHERE scrip_code IN ({placeholders}) ORDER BY date",
            conn,
            params=tuple(symbols),
            parse_dates=["date"],
        )
        volumes = pd.DataFrame(columns=["date", "scrip_code", "volume"])
        if "fact_scrip_volume_history" in tables:
            volumes = pd.read_sql_query(
                "SELECT date, scrip_code, volume FROM fact_scrip_volume_history "
                f"WHERE scrip_code IN ({placeholders}) ORDER BY date",
                conn,
                params=tuple(symbols),
                parse_dates=["date"],
            )
        metadata = pd.DataFrame(columns=["scrip_code", "sector", "free_float_pct", "market_cap_band"])
        if "dim_security_metadata" in tables:
            metadata = pd.read_sql_query(
                "SELECT scrip_code, sector, free_float_pct, market_cap_band "
                "FROM dim_security_metadata",
                conn,
            )
        profiles = pd.DataFrame(columns=["scrip_code", "risk_level", "liquidity_basket"])
        if "fact_portfolio_security_profile" in tables:
            profiles = pd.read_sql_query(
                "SELECT scrip_code, risk_level, liquidity_basket "
                "FROM fact_portfolio_security_profile WHERE fund_name = ?",
                conn,
                params=(fund_name.strip(),),
            )
    except sqlite3.OperationalError as error:
        raise HTTPException(
            status_code=503,
            detail="Portfolio data is not ready. Upload W1, W2, and W3 first.",
        ) from error
    finally:
        conn.close()

    if portfolio.empty:
        raise HTTPException(
            status_code=404,
            detail=f"No portfolio history is available for '{fund_name}'.",
        )
    if market.empty:
        raise HTTPException(
            status_code=503,
            detail="Market price history is unavailable.",
        )

    aligned = (
        portfolio.rename(columns={"price": "portfolio_price"})
        .merge(
            market.rename(columns={"price": "market_price"}),
            on="date",
            how="inner",
        )
        .sort_values("date")
        .set_index("date")
        .dropna()
    )
    returns = aligned.pct_change(fill_method=None).dropna()
    if returns.empty:
        raise HTTPException(
            status_code=422,
            detail="Market price history is unavailable for the selected period.",
        )
    latest_date = returns.index.max()
    selected_month = None
    if period == "1d":
        range_start = latest_date.normalize()
        range_end = range_start
    elif period == "1w":
        range_start = latest_date.normalize() - pd.Timedelta(days=latest_date.weekday())
        range_end = range_start + pd.Timedelta(days=6)
    elif period == "1m":
        selected_month = month or latest_date.strftime("%Y-%m")
        if not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", selected_month):
            raise HTTPException(
                status_code=422,
                detail="Month must use the YYYY-MM format.",
            )
        try:
            month_start = pd.Timestamp(datetime.strptime(selected_month, "%Y-%m"))
        except ValueError as error:
            raise HTTPException(
                status_code=422,
                detail="Month must use the YYYY-MM format.",
            ) from error
        range_start = month_start
        range_end = month_start + pd.DateOffset(months=1) - pd.Timedelta(days=1)
    else:
        quarter_start_month = ((latest_date.month - 1) // 3) * 3 + 1
        range_start = pd.Timestamp(latest_date.year, quarter_start_month, 1)
        range_end = range_start + pd.DateOffset(months=3) - pd.Timedelta(days=1)

    period_returns = returns.loc[
        (returns.index >= range_start) & (returns.index <= range_end)
    ]
    if period_returns.empty:
        raise HTTPException(
            status_code=422,
            detail=(
                f"No overlapping portfolio and market returns are available "
                f"for the selected {period} period."
            ),
        )

    calculation_returns = (
        returns.tail(252) if period == "1d" else period_returns
    )
    window = 252 if period == "1d" else 30
    annualization = math.sqrt(252)
    portfolio_volatility = (
        returns["portfolio_price"]
        .rolling(window, min_periods=window)
        .std()
        * annualization
    )
    market_volatility = (
        returns["market_price"]
        .rolling(window, min_periods=window)
        .std()
        * annualization
    )
    portfolio_std = returns["portfolio_price"].rolling(
        window, min_periods=window
    ).std()
    market_std = returns["market_price"].rolling(
        window, min_periods=window
    ).std()
    portfolio_sharpe = (
        returns["portfolio_price"].rolling(window, min_periods=window).mean()
        / portfolio_std
        * annualization
    )
    market_sharpe = (
        returns["market_price"].rolling(window, min_periods=window).mean()
        / market_std
        * annualization
    )
    risk_metrics = pd.DataFrame(
        {
            "portfolio_volatility": portfolio_volatility,
            "market_volatility": market_volatility,
            "portfolio_sharpe": portfolio_sharpe,
            "market_sharpe": market_sharpe,
        }
    ).dropna()
    if risk_metrics.empty:
        raise HTTPException(
            status_code=422,
            detail="Risk metrics are unavailable for this portfolio.",
        )
    risk_metrics = risk_metrics.loc[
        (risk_metrics.index >= range_start) & (risk_metrics.index <= range_end)
    ]
    if risk_metrics.empty:
        raise HTTPException(
            status_code=422,
            detail="Risk metrics are unavailable for the selected period.",
        )
    risk_metrics.index.name = "date"
    risk_metrics = risk_metrics.reset_index()
    risk_metrics["date"] = pd.to_datetime(risk_metrics["date"]).dt.strftime(
        "%Y-%m-%d"
    )

    security_matrix = pd.DataFrame()
    if not security_prices.empty:
        security_matrix = security_prices.pivot(
            index="date", columns="scrip_code", values="price"
        ).sort_index()
        security_returns = security_matrix.pct_change(fill_method=None)
    else:
        security_returns = pd.DataFrame()

    def finite_or_none(value):
        if value is None or pd.isna(value):
            return None
        number = float(value)
        return number if math.isfinite(number) else None

    snapshot_date = pd.to_datetime(weights["date"].iloc[0])
    volume_averages = {}
    for symbol, symbol_volumes in volumes.groupby("scrip_code"):
        recent = (
            symbol_volumes.loc[symbol_volumes["date"] <= snapshot_date]
            .sort_values("date")
            .tail(20)
        )
        volume_averages[symbol] = (
            pd.to_numeric(recent["volume"], errors="coerce").mean()
            if not recent.empty
            else None
        )
    available_volume_symbols = set(volumes["scrip_code"].dropna().unique())

    metadata_by_symbol = metadata.set_index("scrip_code").to_dict(orient="index")
    profiles_by_symbol = profiles.set_index("scrip_code").to_dict(orient="index")
    market_returns = calculation_returns["market_price"]
    market_variance = market_returns.var(ddof=1)
    holdings_data = []
    risk_by_symbol = {}
    categories = {}
    liquidity_basket_categories = {
        "highly liquid": "Highly Liquid",
        "liquid": "Liquid",
        "semi liquid": "Semi-Liquid",
        "illiquid": "Illiquid",
    }
    for row in weights.to_dict(orient="records"):
        symbol = row["scrip_code"]
        security_return_series = (
            security_returns[symbol]
            if symbol in security_returns
            else pd.Series(dtype=float)
        )
        security_aligned = pd.concat(
            [
                security_return_series.rename("security"),
                market_returns.rename("market"),
            ],
            axis=1,
        ).dropna()
        volatility = security_aligned["security"].std(ddof=1) * annualization
        sharpe = (
            security_aligned["security"].mean()
            / security_aligned["security"].std(ddof=1)
            * annualization
            if security_aligned["security"].std(ddof=1) > 0
            else None
        )
        beta = (
            security_aligned["security"].cov(security_aligned["market"])
            / market_variance
            if market_variance > 0 and len(security_aligned) > 1
            else None
        )
        alpha = (
            (security_aligned["security"].mean() - beta * security_aligned["market"].mean())
            * 252
            if beta is not None and len(security_aligned) > 1
            else None
        )
        details = metadata_by_symbol.get(symbol, {})
        profile = profiles_by_symbol.get(symbol, {})
        average_volume = volume_averages.get(symbol)
        days_to_liquidate = (
            float(row["holding"]) / (float(average_volume) * 0.5)
            if average_volume is not None
            and pd.notna(average_volume)
            and float(average_volume) > 0
            else None
        )
        liquidity_category = profile.get("liquidity_basket")
        if symbol not in available_volume_symbols:
            liquidity_category = "Volume unavailable"
        elif days_to_liquidate is None:
            liquidity_category = "Not traded"
        elif not liquidity_category:
            liquidity_category = "Unclassified"
        category_key = (
            str(liquidity_category).strip().casefold().replace("-", " ")
            if liquidity_category
            else ""
        )
        basket_category = liquidity_basket_categories.get(category_key)
        if basket_category:
            categories.setdefault(basket_category, {"value": 0.0, "days": []})
            categories[basket_category]["value"] += float(row["market_value"])
            if days_to_liquidate is not None:
                categories[basket_category]["days"].append(days_to_liquidate)

        risk_level = profile.get("risk_level")
        risk_by_symbol[symbol] = finite_or_none(volatility)
        holdings_data.append(
            {
                "symbol": symbol,
                "sector": details.get("sector"),
                "market_value": finite_or_none(row["market_value"]),
                "risk_level": risk_level,
                "weight": finite_or_none(row["portfolio_weight"]),
                "volatility": finite_or_none(volatility),
                "market_cap_band": details.get("market_cap_band"),
                "beta": finite_or_none(beta),
                "alpha": finite_or_none(alpha),
                "sharpe": finite_or_none(sharpe),
                "free_float_pct": finite_or_none(details.get("free_float_pct")),
                "liquidity_basket": liquidity_category,
                "average_daily_volume": finite_or_none(average_volume),
                "days_to_liquidate": finite_or_none(days_to_liquidate),
            }
        )

    liquidity = [
        {
            "category": category,
            "value": item["value"],
            "average_days_to_liquidate": (
                sum(item["days"]) / len(item["days"]) if item["days"] else None
            ),
        }
        for category, item in categories.items()
    ]
    liquidity.sort(
        key=lambda item: (
            {
                "Highly Liquid": 0,
                "Liquid": 1,
                "Semi-Liquid": 2,
                "Illiquid": 3,
            }.get(item["category"], 7),
            item["category"],
        )
    )

    portfolio_return_variance = calculation_returns["portfolio_price"].var(ddof=1)
    portfolio_beta = (
        calculation_returns["portfolio_price"].cov(market_returns) / market_variance
        if market_variance > 0 and len(calculation_returns) > 1
        else None
    )
    portfolio_alpha = (
        (
            calculation_returns["portfolio_price"].mean()
            - portfolio_beta * market_returns.mean()
        )
        * 252
        if portfolio_beta is not None
        else None
    )
    latest_metric = risk_metrics.iloc[-1]
    portfolio_var_95 = calculation_returns["portfolio_price"].quantile(0.05)
    portfolio_var_99 = calculation_returns["portfolio_price"].quantile(0.01)
    market_var_99 = market_returns.quantile(0.01)
    summary = {
        "market_value": finite_or_none(weights["market_value"].sum()),
        "market_volatility": finite_or_none(latest_metric["market_volatility"]),
        "portfolio_volatility": finite_or_none(latest_metric["portfolio_volatility"]),
        "alpha": finite_or_none(portfolio_alpha),
        "beta": finite_or_none(portfolio_beta),
        "portfolio_sharpe": finite_or_none(latest_metric["portfolio_sharpe"]),
        "market_sharpe": finite_or_none(latest_metric["market_sharpe"]),
        "market_var_99": finite_or_none(market_var_99),
        "portfolio_var_95": finite_or_none(portfolio_var_95),
        "portfolio_var_99": finite_or_none(portfolio_var_99),
        "days_to_liquidate": finite_or_none(
            sum(
                item["days_to_liquidate"] * item["market_value"]
                for item in holdings_data
                if item["days_to_liquidate"] is not None
                and item["market_value"] is not None
            )
            / sum(
                item["market_value"]
                for item in holdings_data
                if item["days_to_liquidate"] is not None
                and item["market_value"] is not None
            )
            if any(
                item["days_to_liquidate"] is not None
                and item["market_value"] is not None
                for item in holdings_data
            )
            else None
        ),
    }
    if portfolio_return_variance <= 0:
        summary["portfolio_sharpe"] = None
    security_volatility = [
        {"symbol": symbol, "volatility": volatility}
        for symbol, volatility in risk_by_symbol.items()
        if volatility is not None
    ]

    return {
        "fund_name": fund_name,
        "summary": summary,
        "liquidity_basket": liquidity,
        "risk_metrics": risk_metrics.to_dict(orient="records"),
        "security_volatility": security_volatility,
        "holdings": holdings_data,
        "rolling_window": window,
        "period": period,
        "period_start": period_returns.index.min().strftime("%Y-%m-%d"),
        "period_end": period_returns.index.max().strftime("%Y-%m-%d"),
        "as_of_date": latest_date.strftime("%Y-%m-%d"),
        "selected_month": selected_month,
    }


def _equity_chart_export_data(report_data, chart_name):
    if chart_name == "liquidity":
        total = sum(
            float(item["value"])
            for item in report_data["liquidity_basket"]
            if float(item["value"]) > 0
        )
        rows = [
            [
                item["category"],
                float(item["value"]),
                float(item["value"]) / total if total else None,
                item["average_days_to_liquidate"],
            ]
            for item in report_data["liquidity_basket"]
            if float(item["value"]) > 0
        ]
        return (
            "Liquidity Basket",
            ["Liquidity Basket", "Market Value (PKR)", "Share of Basket", "Average Days to Liquidate"],
            rows,
            [],
        )

    if chart_name in {"volatility", "sharpe"}:
        if chart_name == "volatility":
            portfolio_key = "portfolio_volatility"
            market_key = "market_volatility"
            sheet_name = "Volatility"
            portfolio_header = "Portfolio Annualized Volatility"
            market_header = "Market Annualized Volatility"
        else:
            portfolio_key = "portfolio_sharpe"
            market_key = "market_sharpe"
            sheet_name = "Sharpe"
            portfolio_header = "Portfolio Sharpe"
            market_header = "Market Sharpe"
        rows = [
            [item["date"], item[portfolio_key], item[market_key]]
            for item in report_data["risk_metrics"]
        ]
        return (
            sheet_name,
            ["Date", portfolio_header, market_header],
            rows,
            [],
        )

    if chart_name == "distribution":
        securities = report_data["security_volatility"]
        values = sorted(float(item["volatility"]) * 100 for item in securities)
        if not values:
            bins = []
        else:
            minimum = values[0]
            maximum = values[-1]
            bin_count = min(12, max(1, math.ceil(math.sqrt(len(values)))))
            width = (
                max(abs(minimum) * 0.1, 0.01)
                if maximum == minimum
                else (maximum - minimum) / bin_count
            )
            bins = [
                {
                    "lower": minimum + index * width,
                    "upper": minimum + (index + 1) * width,
                    "count": 0,
                }
                for index in range(bin_count)
            ]
            for value in values:
                index = min(
                    bin_count - 1,
                    math.floor((value - minimum) / width),
                )
                bins[index]["count"] += 1
        rows = [
            [
                f'{item["lower"]:.1f}-{item["upper"]:.1f}%',
                item["lower"],
                item["upper"],
                item["count"],
            ]
            for item in bins
        ]
        raw_rows = [
            [item["symbol"], float(item["volatility"]) * 100]
            for item in securities
        ]
        return (
            "Volatility Distribution",
            ["Annualized Volatility Range", "Lower Bound (%)", "Upper Bound (%)", "Number of Symbols"],
            rows,
            [("Constituent Volatility", ["Symbol", "Annualized Volatility (%)"], raw_rows)],
        )

    raise HTTPException(
        status_code=404,
        detail="Chart must be one of: liquidity, volatility, sharpe, distribution.",
    )


def _add_equity_export_sheet(workbook, title, headers, rows):
    worksheet = (
        workbook.active
        if workbook.active.max_row == 1
        and workbook.active.max_column == 1
        and workbook.active["A1"].value is None
        else workbook.create_sheet()
    )
    worksheet.title = title[:31]
    for column_index, header in enumerate(headers, start=1):
        worksheet.cell(row=1, column=column_index, value=header)
    for row in rows:
        worksheet.append(row)
    worksheet.freeze_panes = "A2"
    if headers:
        worksheet.auto_filter.ref = worksheet.dimensions
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="C8102E")
    number_formats = {
        "Market Value (PKR)": "#,##0.00",
        "Share of Basket": "0.00%",
        "Average Days to Liquidate": "0.0",
        "Portfolio Annualized Volatility": "0.00%",
        "Market Annualized Volatility": "0.00%",
        "Lower Bound (%)": '0.0"%"',
        "Upper Bound (%)": '0.0"%"',
        "Annualized Volatility (%)": '0.00"%"',
    }
    for column_index, header in enumerate(headers, start=1):
        number_format = number_formats.get(header)
        if number_format:
            for row_index in range(2, worksheet.max_row + 1):
                worksheet.cell(row=row_index, column=column_index).number_format = number_format
    for column in worksheet.columns:
        letter = column[0].column_letter
        width = min(max(max(len(str(cell.value or "")) for cell in column) + 2, 12), 42)
        worksheet.column_dimensions[letter].width = width
    return worksheet


@app.get("/api/equity-report/{fund_name}/export/{chart_name}")
async def download_equity_chart_data(
    fund_name: str,
    chart_name: str,
    period: str = "1d",
    month: str | None = None,
):
    report_data = await get_equity_report_data(fund_name, period, month)
    sheet_name, headers, rows, additional_sheets = _equity_chart_export_data(
        report_data,
        chart_name,
    )
    workbook = Workbook()
    _add_equity_export_sheet(workbook, sheet_name, headers, rows)
    for title, extra_headers, extra_rows in additional_sheets:
        _add_equity_export_sheet(workbook, title, extra_headers, extra_rows)

    output = BytesIO()
    workbook.save(output)
    safe_fund_name = re.sub(r"[^A-Za-z0-9_-]+", "_", fund_name).strip("_")
    month_suffix = f"_{month}" if period == "1m" and month else ""
    filename = (
        f"{safe_fund_name or 'portfolio'}_{chart_name}_{period}{month_suffix}.xlsx"
    )
    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )